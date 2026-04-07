"""生成：测试用例表_手动全量.xlsx（中文手动测试，含实际结果/状态列）。"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = [
    "序号",
    "模块",
    "测试点",
    "前置条件",
    "操作步骤",
    "预期结果",
    "实际结果",
    "状态",
    "备注",
]

# 每项：(模块, 测试点, 前置条件, 操作步骤, 预期结果)
ROWS: list[tuple[str, str, str, str, str]] = [
    ("环境与部署", "本地启动", "已 clone；已安装 Node 与 pnpm", "在项目根目录执行 pnpm install 后 pnpm dev", "终端无致命错误，浏览器可访问本地 URL"),
    ("环境与部署", "环境变量", "存在 .env.local", "核对 NEXT_PUBLIC_SUPABASE_URL、密钥等配置与团队文档一致", "应用可连 Supabase；无关键变量缺失导致的启动失败"),
    ("环境与部署", "数据库", "可登录 Supabase 控制台", "按仓库说明执行 schema.sql、seed.sql（或团队约定迁移）", "首页与列表能加载优惠等基础数据"),
    ("首页 /", "整体渲染", "服务已启动", "打开 /，从上到下浏览各区块", "Hero、Features、Offres、HowToUse、Stats、Testimonials 正常显示，无白屏"),
    ("首页 /", "导航栏", "同左", "点击 Navbar 中 Offres、Lieux、Cadeaux、Activités、Profil 等", "跳转路径正确；未登录访问 Profil 应进入登录流程"),
    ("首页 /", "页脚", "同左", "点击页脚各链接", "合法链接可打开；已知未实现项记录备注"),
    ("首页 /", "深色模式", "同左", "使用 ThemeToggle 切换浅色/深色", "主题切换正常，对比度可接受"),
    ("首页 /", "Hero 搜索", "同左", "在 Hero 搜索框输入关键词并触发搜索（按页面按钮）", "进入 /offres 且搜索结果与关键词一致（或 URL 含 q=）"),
    ("首页 /", "Hero 地区", "Hero 有地区建议", "选择某地区并触发跳转列表", "进入 /offres 且地区过滤与所选一致（含 loc 参数时）"),
    ("首页 /", "Offres 分类 Tab", "同左", "在首页 Offres 区块切换 Tout、Restaurants 等", "卡片列表随分类变化"),
    ("首页 /", "进入详情", "同左", "点击某张 Offer 卡片", "进入 /offres/[id]，信息与该优惠一致"),
    ("首页 /", "距离（未定位）", "浏览器未授予定位权限", "不点击「Ma position」，查看卡片", "卡片不展示距离（distance 为空不显示）"),
    ("首页 /", "距离（已定位）", "浏览器可定位", "点击「Ma position」并允许权限", "卡片显示与当前位置相关的 m 或 km"),
    ("首页 /", "定位拒绝", "—", "点击定位并拒绝浏览器权限", "出现拒绝/重试类提示，页面不崩溃"),
    ("列表 /offres", "重定向 cadeaux", "服务正常", "浏览器打开 /cadeaux", "重定向到 /offres?cat=cadeau 且分类为 Cadeaux"),
    ("列表 /offres", "重定向 activites", "服务正常", "浏览器打开 /activites", "重定向到 /offres?cat=activite 且分类正确"),
    ("列表 /offres", "URL 初始分类", "—", "直接访问 /offres?cat=restaurant（或其它合法 cat）", "对应 Tab 选中，列表仅该类"),
    ("列表 /offres", "搜索过滤", "—", "在搜索框输入标题或描述中的词", "列表仅保留匹配项"),
    ("列表 /offres", "地区参数 loc", "—", "访问 /offres?loc=某城市或邮编片段", "页面上方显示地区说明；列表与地址/标题匹配"),
    ("列表 /offres", "距离（未定位）", "未授权定位", "不点「Ma position」", "卡片不显示距离"),
    ("列表 /offres", "距离与排序", "已授权定位", "点击「Ma position」直至成功", "显示距离；列表按距离由近到远；无坐标条目顺序合理"),
    ("列表 /offres", "定位状态 UI", "—", "观察定位按钮：空闲/加载/成功/拒绝/错误", "文案与样式与状态一致；拒绝时有简短引导"),
    ("详情 /offres/[id]", "正常展示", "存在有效优惠 id", "打开一条数据完整的优惠", "标题、描述、价格、地址、详情列表等正确"),
    ("详情 /offres/[id]", "404", "—", "访问不存在的 id", "显示项目 404 页"),
    ("详情 /offres/[id]", "Leaflet 地图", "优惠有坐标或可 geocode 的地址", "滚动到地图区域并等待加载", "地图与标记显示；控制台无地图初始化错误"),
    ("详情 /offres/[id]", "无坐标无地址", "库中存在此类数据", "打开对应优惠", "显示无法展示地图的占位文案，不崩溃"),
    ("详情 /offres/[id]", "Google Maps 链接", "有地址或坐标", "点击打开 Google 地图", "新标签页打开且查询位置合理"),
    ("详情 /offres/[id]", "收藏（未登录）", "已登出", "打开详情页查看收藏按钮", "不崩溃；未登录不表现为已收藏（401 处理正常）"),
    ("详情 /offres/[id]", "收藏（添加）", "已登录", "点击加入收藏", "成功；UI 为已收藏；刷新后仍一致"),
    ("详情 /offres/[id]", "收藏（取消）", "已登录且已收藏", "再次点击取消收藏", "成功；状态恢复且刷新一致"),
    ("详情 /offres/[id]", "评价列表", "该优惠在库中有评价", "查看页面评价区块", "展示 Supabase 中的评价"),
    ("详情 /offres/[id]", "评价（未登录）", "已登出", "尝试填写并提交评价", "被阻止或提示登录"),
    ("详情 /offres/[id]", "评价（提交）", "已登录", "填写评分、标题、正文并提交", "成功反馈；列表出现新评价或刷新后可见"),
    ("详情 /offres/[id]", "相关优惠", "同左", "查看页底相关推荐", "展示同类别其它条目，链接可点"),
    ("Profil /profil", "未登录拦截", "已登出", "直接访问 /profil", "重定向到 /connexion（或约定登录页）"),
    ("Profil /profil", "收藏展示", "已登录且 favorites 表有数据", "打开个人页", "展示最近收藏，与数据库一致"),
    ("Profil /profil", "我的评价", "已登录且曾提交评价", "查看个人页评价区域", "列出当前用户的评价记录"),
    ("Profil /profil", "资料保存", "已登录", "修改个人资料并保存（若有）", "请求成功，刷新后信息更新"),
    ("Profil /profil", "删除账号", "使用测试小号", "执行删除账号（若页面提供）", "账号无法再登录；会话清除（按产品设计）"),
    ("注册 /inscription", "成功注册", "使用未注册邮箱", "完整填写表单并提交", "注册成功或进入下一步；失败时有明确提示"),
    ("注册 /inscription", "异常输入", "—", "重复邮箱、弱密码、空字段等", "前端或接口返回合理错误"),
    ("登录 /connexion", "成功登录", "账号已存在", "输入正确邮箱与密码", "登录成功；可访问 /profil"),
    ("登录 /connexion", "错误密码", "—", "故意输错密码", "失败提示；不处于已登录态"),
    ("会话", "登出", "已登录", "在界面执行登出", "再访问 /profil 需重新登录"),
    ("Lieux /lieux", "地图与列表", "服务与数据正常", "打开 /lieux，操作地图、列表、定位等（按实际 UI）", "行为与数据一致，无严重控制台错误"),
    ("Admin /admin", "未授权（可选）", "非管理员或未登录", "直接访问 /admin", "跳转登录、403 或拒绝访问（与项目策略一致）"),
    ("Admin /admin", "管理员（可选）", "具备管理员权限", "登录后访问 admin 子页面", "仅授权用户可操作；无越权"),
    ("黑盒 API", "收藏 GET", "已登录；打开开发者工具 Network", "打开含收藏按钮的详情页", "/api/favorites?offerId= 返回 200 与 isFavorite 正确"),
    ("黑盒 API", "收藏 POST/DELETE", "已登录", "切换收藏开/关", "对应 POST/DELETE 为 2xx；失败时有响应或前端提示"),
    ("黑盒 API", "评价 POST", "已登录", "提交一条评价", "/api/reviews 或等价路由为 2xx；列表可见"),
    ("回归", "与 CI 一致", "本地环境", "运行 pnpm lint、pnpm build（或 package.json 中与 CI 相同脚本）", "与 GitHub Actions 结果一致通过"),
    ("回归", "基础体验", "—", "多次刷新、快速切换路由", "无明显卡顿、无未捕获错误弹窗"),
]


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    out = root / "测试用例表_手动全量.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "手动测试"

    header_fill = PatternFill("solid", fgColor="FF4472C4")
    header_font = Font(bold=True, color="FFFFFFFF")

    for col, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, (m, p, pre, step, expect) in enumerate(ROWS, start=1):
        r = i + 1
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=2, value=m)
        ws.cell(row=r, column=3, value=p)
        ws.cell(row=r, column=4, value=pre)
        ws.cell(row=r, column=5, value=step)
        ws.cell(row=r, column=6, value=expect)
        ws.cell(row=r, column=7, value="")
        ws.cell(row=r, column=8, value="未测")
        ws.cell(row=r, column=9, value="")
        for c in range(1, 10):
            al = ws.cell(row=r, column=c).alignment
            ws.cell(row=r, column=c).alignment = Alignment(
                horizontal=al.horizontal if c == 8 else "left",
                vertical="top",
                wrap_text=True,
            )

    widths = [6, 16, 22, 28, 40, 40, 24, 10, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    wb.save(out)
    print(f"Wrote {out} ({len(ROWS)} cases)")


if __name__ == "__main__":
    main()
